"""Gera docs/base_conhecimento.xlsx para o grupo revisar a base em planilha."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

RAIZ = Path(__file__).resolve().parent.parent
base = json.loads((RAIZ / "base" / "base_conhecimento.json").read_text(encoding="utf-8"))

wb = Workbook()
ws = wb.active
ws.title = "Registros"
cab = ["id", "tema", "perguntas", "resumo", "resposta", "aprofundamento", "fontes", "limite", "acao", "seguranca"]
ws.append(cab)
regs = base["registros"] + list(base["respostas_seguranca"].values())
for r in regs:
    ws.append([r["id"], r["tema"], "\n".join(r.get("perguntas", [])), r["resumo"], r["resposta"], r["aprofundamento"],
               "\n".join(base["fontes"][f] for f in r["fontes"]), r["limite"], r["acao"] or "", r["seguranca"] or ""])
for col, larg in zip("ABCDEFGHIJ", [22, 22, 40, 45, 60, 70, 50, 35, 20, 20]):
    ws.column_dimensions[col].width = larg
for row in ws.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
for c in ws[1]:
    c.font = Font(bold=True)
ws.freeze_panes = "A2"

wg = wb.create_sheet("Glossário")
wg.append(["termo", "formas reconhecidas", "definição"])
for g in base["glossario"]:
    wg.append([g["termo"], ", ".join(g["formas"]), g["definicao"]])
for col, larg in zip("ABC", [28, 40, 90]):
    wg.column_dimensions[col].width = larg

wf = wb.create_sheet("Fontes")
wf.append(["chave", "referência"])
for k, v in base["fontes"].items():
    wf.append([k, v])
wf.column_dimensions["A"].width = 18
wf.column_dimensions["B"].width = 120

wb.save(RAIZ / "docs" / "base_conhecimento.xlsx")
